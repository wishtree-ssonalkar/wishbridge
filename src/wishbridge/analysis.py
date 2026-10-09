"""Step 1 - Analyze: run the LakeBridge analyzer and turn its Excel report into structured data + an effort estimate."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from . import lakebridge
from .staging import staged
from .config import ProjectConfig
from .state import save_step


def _table_rows(ws, first_header: str) -> list[dict[str, Any]]:
    """Rows of a sheet as dicts, using the first row that contains `first_header` as the header."""
    rows = list(ws.iter_rows(values_only=True))
    idx = next((i for i, r in enumerate(rows) if r and first_header in r), None)
    if idx is None:
        return []
    header = [str(h) if h is not None else "" for h in rows[idx]]
    key = header.index(first_header)
    return [dict(zip(header, r)) for r in rows[idx + 1:] if r and len(r) > key and r[key] not in (None, "")]


def parse_report(xlsx: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    totals: dict[str, int] = {}
    programs: list[dict[str, Any]] = []
    functions: dict[str, int] = {}

    if "Summary" in wb.sheetnames:
        for row in wb["Summary"].iter_rows(values_only=True):
            cells = [c for c in row if c not in (None, "")]
            for i, c in enumerate(cells[:-1]):
                if isinstance(c, str) and c.startswith("Total ") and isinstance(cells[i + 1], (int, float)):
                    totals[c.removeprefix("Total ").strip()] = int(cells[i + 1])

    if "SQL Programs" in wb.sheetnames:
        rows = list(wb["SQL Programs"].iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows) if r and "Program Name" in r), None)
        if header_idx is not None:
            header = [str(h) if h is not None else "" for h in rows[header_idx]]
            col = {h: i for i, h in enumerate(header)}
            for r in rows[header_idx + 1:]:
                if not r or r[col["Program Name"]] in (None, ""):
                    continue

                def get(name: str, default: Any = None) -> Any:
                    i = col.get(name)
                    return r[i] if i is not None and i < len(r) and r[i] is not None else default

                programs.append({
                    "name": str(get("Program Name")),
                    "source_file": str(get("Source File", "")),
                    "lines": int(get("Line Count", 0) or 0),
                    "statements": int(get("Statement Count", 0) or 0),
                    "complexity": str(get("Complexity", "UNKNOWN")).upper(),
                    "category": str(get("Script Category", "")),
                })

    # ETL sources (SSIS, Informatica, DataStage): one row per job / package instead of per SQL program.
    if not programs and "Job Details" in wb.sheetnames:
        for job in _table_rows(wb["Job Details"], "Job Name"):
            if str(job.get("Included") or "YES").upper() == "NO":
                continue
            programs.append({
                "name": str(job["Job Name"]),
                "source_file": str(job.get("Source File") or ""),
                "lines": 0,
                "statements": int(job.get("Number of Nodes") or 0),
                "complexity": str(job.get("Categorization") or "UNKNOWN").upper(),
                "category": str(job.get("Job Type") or "Job"),
            })
    if "Embedded SQL Programs" in wb.sheetnames:
        embedded = _table_rows(wb["Embedded SQL Programs"], "Program Name")
        if embedded:
            totals["Embedded SQL statements"] = len(embedded)

    if "Functions" in wb.sheetnames:
        for r in list(wb["Functions"].iter_rows(values_only=True))[1:]:
            if r and r[0] and isinstance(r[1], (int, float)):
                functions[str(r[0])] = int(r[1])

    wb.close()
    complexity: dict[str, int] = {}
    for p in programs:
        complexity[p["complexity"]] = complexity.get(p["complexity"], 0) + 1
    return {"totals": totals, "programs": programs, "functions": functions, "complexity": complexity}


def estimate_hours(cfg: ProjectConfig, programs: list[dict[str, Any]]) -> float:
    return round(sum(cfg.hours_per_file.get(p["complexity"], cfg.hours_per_file["MEDIUM"]) for p in programs), 1)


def run_analyze(cfg: ProjectConfig) -> dict[str, Any]:
    if not cfg.input_dir.exists() or not any(cfg.input_dir.rglob("*")):
        raise FileNotFoundError(f"No input files in {cfg.input_dir}")
    report = cfg.out("analysis", "analysis.xlsx")
    lakebridge.analyze(staged(cfg), report)  # code files only: no bin/obj or project files
    if not report.exists():
        # Some analyzer versions drop the extension.
        bare = report.with_suffix("")
        if bare.exists():
            bare.replace(report)
    result = parse_report(report)
    result["report_file"] = str(report)
    result["estimated_hours_baseline"] = estimate_hours(cfg, result["programs"])
    save_step(cfg, "analyze", result)
    from .fit import run_fit

    fit = run_fit(cfg)
    result["fit"] = {"verdict": fit["verdict"], "headline": fit["headline"], "counts": fit["counts"]}
    return result
