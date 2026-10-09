"""SSIS: master-package orchestration, notebook checks, ETL analyzer workbooks."""

from pathlib import Path

import openpyxl

from wishbridge.analysis import parse_report
from wishbridge.config import load_config
from wishbridge.convert import orchestration_findings
from wishbridge.orchestration import job_tasks, ssis_plan, write_job_definition
from wishbridge.rules import check_notebook, detect

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "ssis-demo"


def test_master_package_order_becomes_job_dependencies():
    cfg = load_config(EXAMPLE / "project.yml")
    plan = ssis_plan(cfg)
    assert plan["orchestrators"] == {"Master": ["LoadDimCustomer", "LoadFactSales"]}
    assert plan["depends"] == {"LoadDimCustomer": [], "LoadFactSales": ["LoadDimCustomer"]}
    notebooks = [{"file": f"SalesDW_ETL/{n}.py", "path": f"/Workspace/x/SalesDW_ETL/{n}"}
                 for n in ("LoadDimCustomer", "LoadFactSales", "Master")]
    tasks, skipped = job_tasks(cfg, notebooks)
    assert [(t["key"], t["depends_on"]) for t in tasks] == [("LoadDimCustomer", []), ("LoadFactSales", ["LoadDimCustomer"])]
    assert skipped == {"SalesDW_ETL/Master.py": ["LoadDimCustomer", "LoadFactSales"]}


def test_job_definition_can_be_scheduled(tmp_path):
    cfg = load_config(EXAMPLE / "project.yml")
    from dataclasses import replace
    import json

    cfg = replace(cfg, output_dir=tmp_path)
    out = write_job_definition(cfg, [{"key": "A", "path": "/p/A", "depends_on": []},
                                     {"key": "B", "path": "/p/B", "depends_on": ["A"]}])
    job = json.loads(out.read_text(encoding="utf-8"))
    assert job["tasks"][1] == {"task_key": "B", "notebook_task": {"notebook_path": "/p/B"}, "depends_on": [{"task_key": "A"}]}


def test_empty_master_notebook_is_explained_and_other_empty_notebooks_warned():
    empty = "# Databricks notebook source\nimport sys\nPackageName = 'Master'\n"
    info = orchestration_findings(Path("SalesDW_ETL/Master.py"), empty, {"Master": ["A", "B"]})
    assert info[0].rule == "orchestration" and "A -> B" in info[0].message
    warn = orchestration_findings(Path("Other.py"), empty, {})
    assert warn[0].rule == "empty-notebook" and warn[0].severity == "warning"


def test_ssis_expressions_in_notebook_sql_are_fixed_and_flagged():
    nb = ('# Databricks notebook source\nx = "a" + b  # python stays as it is\n'
          'q = f"""SELECT TRIM(FirstName) + ` ` + TRIM(LastName) AS FullName,\n'
          'LOWER(SUBSTRING(Email, FINDSTRING(Email, `@`, 1) + 1, LEN(Email))) AS Domain,\n'
          '(DT_WSTR, 10) Code, `Order Date`\nFROM Sales.Customer WHERE c.IsActive = 1"""\n'
          'a = f"""SELECT\n\nFROM t"""\n')
    fixed, findings = check_notebook(nb, {"Sales": "main.src"})
    assert 'x = "a" + b' in fixed
    assert "TRIM(FirstName) || ' ' || TRIM(LastName)" in fixed
    assert "instr(Email, '@')" in fixed and "length(Email)" in fixed
    assert "`Order Date`" in fixed and "FROM main.src.Customer" in fixed
    open_rules = {f.rule for f in findings if not f.fixed}
    assert {"empty-select", "etl-cast", "bit-flag-compare"} <= open_rules
    assert {"etl-quoted-literal", "etl-string-concat", "etl-findstring", "etl-len", "schema-map"} <= \
        {f.rule for f in findings if f.fixed}


def test_bit_flag_compare_in_tsql():
    rules = {f.rule for f in detect("SELECT * FROM u WHERE u.IsDeleted = 0 AND Status = 1;", "mssql")}
    assert "bit-flag-compare" in rules
    assert "bit-flag-compare" not in {f.rule for f in detect("SELECT * FROM u WHERE Status = 1;", "mssql")}


def test_etl_analyzer_workbook_lists_jobs(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Total Jobs", 0])
    jd = wb.create_sheet("Job Details")
    jd.append(["Job Name", "Folder", "Source File", "Included", "Job Type", "Categorization", "Number of Nodes"])
    jd.append(["P.LoadA", "f", "f/LoadA.dtsx", "YES", "Package", "LOW", 6])
    jd.append(["P.LoadB", "f", "f/LoadB.dtsx", "YES", "Package", "HIGH", 20])
    es = wb.create_sheet("Embedded SQL Programs")
    es.append(["Program Name", "Source File"])
    es.append(["a.sql", "x"])
    out = tmp_path / "a.xlsx"
    wb.save(out)
    r = parse_report(out)
    assert [p["name"] for p in r["programs"]] == ["P.LoadA", "P.LoadB"]
    assert r["complexity"] == {"LOW": 1, "HIGH": 1}
    assert r["totals"]["Embedded SQL statements"] == 1


def test_etl_export_formats_are_detected_and_explained(tmp_path):
    import zipfile

    import pytest

    from wishbridge.config import ConfigError, load_config
    from wishbridge.discover import detect_source
    from wishbridge.staging import staged

    ds = tmp_path / "ds"
    ds.mkdir()
    (ds / "j.xml").write_text('<?xml version="1.0"?>\n<DSExport><Header/></DSExport>\n', encoding="utf-8")
    assert detect_source(ds).source == "datastage"

    iics = tmp_path / "iics"
    iics.mkdir()
    with zipfile.ZipFile(iics / "m.zip", "w") as z:
        z.writestr("exportMetadata.v2.json", "{}")
    assert detect_source(iics).source == "informatica-cloud"

    def project(name, source, files):
        code = tmp_path / name
        code.mkdir()
        for f in files:
            (code / f).write_text("x", encoding="utf-8")
        p = tmp_path / f"{name}.yml"
        p.write_text(f"source: {source}\ninput: '{code.as_posix()}'\noutput: '{(tmp_path / (name + '_out')).as_posix()}'\n",
                     encoding="utf-8")
        return load_config(p)

    with pytest.raises(ConfigError, match="exported as XML"):
        staged(project("dsx", "datastage", ["j.dsx"]))
    with pytest.raises(ConfigError, match="export packages"):
        staged(project("infaxml", "informatica-cloud", ["m.xml"]))
    cloud = load_config(tmp_path / "infaxml.yml")
    assert cloud.target_technology == "PYSPARK"  # BladeBridge generates Informatica Cloud only as PySpark
